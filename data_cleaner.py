"""
data_cleaner.py
===============
Parser robusto para los archivos Excel del BCE (III2025.zip).

Maneja:
  - Encabezados multinivel con celdas fusionadas (openpyxl)
  - Detección automática de la fila de inicio de datos ("Nacional")
  - Separación de niveles geográficos: Nacional / Provincial / Cantonal
  - Hojas con estructura de serie de tiempo (IncFinEmpresas, etc.)
  - Hojas con estructura plana (tablas simples)
"""

from __future__ import annotations

import re
import sqlite3
import logging
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)


# ─── Utilidades generales ────────────────────────────────────────────────────

def _sanitize(name: str) -> str:
    """Convierte texto a snake_case apto para SQL."""
    s = str(name).strip()
    s = re.sub(r"[\s\-/\\().,;:]+", "_", s)
    s = re.sub(r"[^\w]", "", s)
    s = re.sub(r"_+", "_", s).strip("_")
    return s.lower() or "col"


def _is_empty(val) -> bool:
    if val is None:
        return True
    if isinstance(val, float) and np.isnan(val):
        return True
    return str(val).strip() in ("", "nan", "None")


def _coerce_num(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


# ─── Lector raw con openpyxl (maneja celdas fusionadas) ─────────────────────

class RawSheetReader:
    """
    Lee una hoja de Excel con openpyxl y expande celdas fusionadas
    (forward-fill horizontal y vertical).
    """

    def __init__(self, wb_path: Path, sheet_name: str):
        wb = openpyxl.load_workbook(wb_path, data_only=True, read_only=False)
        ws = wb[sheet_name]

        # Primero expandir rangos fusionados
        self._expand_merges(ws)

        # Leer todas las celdas en una matriz
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))

        self.raw: list[list] = rows
        self.n_rows = len(rows)
        self.n_cols = max((len(r) for r in rows), default=0)
        wb.close()

    @staticmethod
    def _expand_merges(ws) -> None:
        """Copia el valor de la celda origen a todas las celdas fusionadas."""
        # Necesitamos copiar los rangos fusionados antes de unmerge
        merge_ranges = list(ws.merged_cells.ranges)
        for mr in merge_ranges:
            top_left_value = ws.cell(mr.min_row, mr.min_col).value
            ws.unmerge_cells(str(mr))
            for row in range(mr.min_row, mr.max_row + 1):
                for col in range(mr.min_col, mr.max_col + 1):
                    ws.cell(row, col).value = top_left_value

    def as_dataframe(self) -> pd.DataFrame:
        """Devuelve la hoja completa como DataFrame sin encabezado."""
        return pd.DataFrame(self.raw)

    def row_as_strings(self, idx: int) -> list[str]:
        if idx >= self.n_rows:
            return []
        return [str(v).strip() if v is not None else "" for v in self.raw[idx]]

    def count_non_empty(self, row_idx: int) -> int:
        return sum(1 for v in self.raw[row_idx] if not _is_empty(v))


# ─── Detección del tipo de estructura ────────────────────────────────────────

NACIONAL_KEYWORDS = re.compile(
    r"\bnacional\b|\btotal\s+nacional\b|\bnacional\s+total\b",
    re.IGNORECASE
)
PROVINCIAL_KEYWORDS = re.compile(
    r"\bprovincial\b|\bprovincia\b|\bprovincias\b",
    re.IGNORECASE
)
CANTONAL_KEYWORDS = re.compile(
    r"\bcantonal\b|\bcant[oó]n\b|\bcantoness?\b",
    re.IGNORECASE
)
INDICATOR_KEYWORDS = re.compile(
    r"\bindicador\b|\buso\b|\bacceso\b|\bcaptaci[oó]n\b|\bcolocaci[oó]n\b",
    re.IGNORECASE
)


def detect_structure(reader: RawSheetReader) -> str:
    """
    Detecta si la hoja tiene estructura:
    - 'timeseries'  → encabezados multinivel + secciones geográficas en filas
    - 'flat'        → tabla plana simple
    """
    full_text = " ".join(
        str(v) for row in reader.raw[:30] for v in row if not _is_empty(v)
    )
    if INDICATOR_KEYWORDS.search(full_text) and NACIONAL_KEYWORDS.search(full_text):
        return "timeseries"
    return "flat"


# ─── Parser de encabezados multinivel ────────────────────────────────────────

def find_header_block(reader: RawSheetReader, max_scan: int = 40) -> tuple[int, int]:
    """
    Devuelve (header_start_row, data_start_row).
    header_start_row = primera fila con suficientes valores no vacíos.
    data_start_row   = fila donde aparece "Nacional" o primer número.
    """
    header_start = None
    for i in range(min(max_scan, reader.n_rows)):
        if reader.count_non_empty(i) >= 3:
            header_start = i
            break

    if header_start is None:
        return 0, 1

    # Buscar el inicio de datos: fila con "Nacional" o primera fila numérica
    data_start = header_start + 1
    for i in range(header_start + 1, min(header_start + 20, reader.n_rows)):
        row_str = " ".join(reader.row_as_strings(i))
        if NACIONAL_KEYWORDS.search(row_str):
            data_start = i
            break
        # Fila predominantemente numérica
        vals = [v for v in reader.raw[i] if not _is_empty(v)]
        num_count = sum(1 for v in vals if isinstance(v, (int, float)))
        if vals and num_count / len(vals) > 0.5:
            data_start = i
            break

    return header_start, data_start


def build_column_names(reader: RawSheetReader,
                        header_start: int,
                        data_start: int) -> list[str]:
    """
    Construye nombres de columna combinando los niveles de encabezado.
    Maneja forward-fill horizontal (celdas fusionadas ya expandidas).
    """
    n_header_rows = data_start - header_start
    if n_header_rows <= 0:
        n_header_rows = 1

    # Recopilar filas de encabezado como strings
    header_rows: list[list[str]] = []
    for r in range(header_start, header_start + n_header_rows):
        row_strs = reader.row_as_strings(r)
        # Asegurar longitud uniforme
        while len(row_strs) < reader.n_cols:
            row_strs.append("")
        header_rows.append(row_strs)

    if not header_rows:
        return [f"col_{i}" for i in range(reader.n_cols)]

    # Abreviaciones para nombres de columna muy largos
    _ABBREV = {
        "indicador uso": "ind",
        "número de clientes": "n_clientes",
        "número de cuentas": "n_cuentas",
        "productos financieros activos": "prod_activos",
        "productos financieros": "prod_fin",
        "empresas con algún producto financiero total": "total",
        "empresas con algún producto financiero de captación": "captacion",
        "empresas con algún producto financiero de colocación": "colocacion",
        "personas con algún producto financiero total": "total",
        "personas con algún producto financiero de captación": "captacion",
        "personas con algún producto financiero de colocación": "colocacion",
    }

    def _abbrev(text: str) -> str:
        t = text.lower().strip()
        for pattern, repl in _ABBREV.items():
            t = t.replace(pattern, repl)
        return t

    # Combinar niveles verticalmente para cada columna
    col_names = []
    seen: dict[str, int] = {}
    for col_idx in range(reader.n_cols):
        parts = []
        for hr in header_rows:
            val = hr[col_idx] if col_idx < len(hr) else ""
            val_clean = val.strip()
            if val_clean and val_clean.lower() not in ("nan", "none"):
                # Evitar duplicar el mismo texto en niveles consecutivos
                if not parts or parts[-1] != val_clean:
                    parts.append(_abbrev(val_clean))
        combined = "__".join(parts) if parts else f"col_{col_idx}"
        sanitized = _sanitize(combined)[:60]  # máx 60 caracteres

        # Desambiguar duplicados
        if sanitized in seen:
            seen[sanitized] += 1
            sanitized = f"{sanitized}_{seen[sanitized]}"
        else:
            seen[sanitized] = 0
        col_names.append(sanitized)

    return col_names


# ─── Parser de secciones geográficas ─────────────────────────────────────────

ECUADOR_PROVINCES = {
    "azuay", "bolivar", "cañar", "carchi", "chimborazo", "cotopaxi",
    "el oro", "esmeraldas", "galapagos", "galápagos", "guayas", "imbabura",
    "loja", "los rios", "los ríos", "manabi", "manabí", "morona santiago",
    "napo", "orellana", "pastaza", "pichincha", "santa elena",
    "santo domingo", "sucumbios", "sucumbíos", "tungurahua", "zamora chinchipe",
}


def classify_geo_row(cell_value: str) -> str:
    """Clasifica un valor de fila como 'nacional', 'provincial', 'cantonal' o 'other'."""
    v = str(cell_value).strip().lower()
    if not v or v in ("nan", "none", ""):
        return "other"
    if NACIONAL_KEYWORDS.search(v):
        return "nacional"
    if any(p in v for p in ECUADOR_PROVINCES):
        return "provincial"
    if PROVINCIAL_KEYWORDS.search(v):
        return "provincial"
    if CANTONAL_KEYWORDS.search(v):
        return "cantonal"
    return "other"


# ─── Parser principal de hojas ────────────────────────────────────────────────

def parse_sheet_timeseries(
    wb_path: Path,
    sheet_name: str,
) -> dict[str, pd.DataFrame]:
    """
    Parsea una hoja con estructura de serie de tiempo.
    Devuelve dict con claves 'nacional', 'provincial', 'cantonal'.
    """
    reader = RawSheetReader(wb_path, sheet_name)
    header_start, data_start = find_header_block(reader)
    col_names = build_column_names(reader, header_start, data_start)

    # Leer datos desde data_start
    data_rows = reader.raw[data_start:]
    if not data_rows:
        return {}

    # Ajustar ancho de columnas
    max_data_cols = max(len(r) for r in data_rows)
    while len(col_names) < max_data_cols:
        col_names.append(f"col_{len(col_names)}")

    df = pd.DataFrame(data_rows, columns=col_names[:max_data_cols])
    df = df.dropna(how="all")

    # La primera columna no vacía suele ser la etiqueta de fila (fecha/geo)
    # Identificar columna de etiqueta (primera con texto)
    label_col = df.columns[0]
    for c in df.columns:
        non_null = df[c].dropna()
        if len(non_null) > 0 and not all(
            isinstance(v, (int, float)) or
            (isinstance(v, str) and re.match(r"^\d+\.?\d*$", v.strip()))
            for v in non_null
        ):
            label_col = c
            break

    # Determinar si hay una columna de período/fecha
    # (buscar columna con valores tipo trimestre/año)
    period_col = None
    for c in df.columns[:5]:
        sample = df[c].dropna().head(10).astype(str)
        if sample.str.match(r".*\d{4}.*").mean() > 0.3:
            period_col = c
            break

    # Separar secciones geográficas
    results: dict[str, list] = {"nacional": [], "provincial": [], "cantonal": [], "other": []}

    current_geo = "nacional"   # Por defecto la primera sección es nacional

    for _, row in df.iterrows():
        # Recopilar todos los valores de texto de la fila
        text_vals = [
            str(v).strip() for v in row
            if not _is_empty(v) and isinstance(v, str) and len(str(v).strip()) > 1
        ]

        if not text_vals:
            continue

        # Contar valores numéricos en la fila
        num_vals = sum(
            1 for v in row
            if isinstance(v, (int, float))
            and not isinstance(v, bool)
            and not np.isnan(float(v))
        )

        # ── Buscar clasificación geográfica en TODOS los textos de la fila ──
        geo_from_row = "other"
        best_text_for_label = text_vals[0] if text_vals else ""

        for tv in text_vals:
            g = classify_geo_row(tv)
            if g in ("nacional", "provincial", "cantonal"):
                geo_from_row = g
                best_text_for_label = tv
                break

        # ── Fila de sólo texto → es un marcador de sección ──
        if num_vals == 0:
            if geo_from_row in ("nacional", "provincial", "cantonal"):
                current_geo = geo_from_row
            # Ignorar sub-encabezados de texto puro
            continue

        # ── Fila con datos numéricos ──
        # El nivel geográfico es el que se detectó en la fila,
        # o si no hay señal clara, el nivel de la sección actual.
        if geo_from_row in ("nacional", "provincial", "cantonal"):
            category = geo_from_row
        else:
            category = current_geo

        row_dict = row.to_dict()
        row_dict["_geo_label"] = best_text_for_label
        row_dict["_geo_level"] = category
        results[category].append(row_dict)

    dfs = {}
    for key, rows in results.items():
        if rows:
            d = pd.DataFrame(rows)
            d = d.dropna(how="all")
            dfs[key] = d

    return dfs


def parse_sheet_flat(wb_path: Path, sheet_name: str) -> pd.DataFrame | None:
    """
    Parsea una hoja con estructura plana.
    """
    reader = RawSheetReader(wb_path, sheet_name)
    header_start, data_start = find_header_block(reader, max_scan=60)
    col_names = build_column_names(reader, header_start, data_start)

    data_rows = reader.raw[data_start:]
    if not data_rows:
        return None

    max_cols = max(len(r) for r in data_rows)
    while len(col_names) < max_cols:
        col_names.append(f"col_{len(col_names)}")

    df = pd.DataFrame(data_rows, columns=col_names[:max_cols])
    df = df.dropna(how="all").dropna(axis=1, how="all")

    threshold = max(1, len(df.columns) // 2)
    df = df.dropna(thresh=threshold)

    return df if not df.empty else None


# ─── Procesamiento completo de un Excel ──────────────────────────────────────

def process_excel_to_sqlite(
    wb_path: Path,
    db_path: Path,
    progress_callback=None,
) -> list[dict]:
    """
    Procesa todos las hojas de un Excel y las escribe en SQLite.
    Devuelve el manifiesto de tablas creadas.
    """
    db_path.parent.mkdir(parents=True, exist_ok=True)
    manifest = []

    try:
        xl = pd.ExcelFile(wb_path, engine="openpyxl" if wb_path.suffix == ".xlsx" else "xlrd")
        sheet_names = xl.sheet_names
    except Exception as e:
        log.error("No se pudo abrir %s: %s", wb_path.name, e)
        return manifest

    log.info("Procesando '%s' (%d hojas)…", wb_path.name, len(sheet_names))

    with sqlite3.connect(db_path) as conn:
        for i, sheet in enumerate(sheet_names):
            if progress_callback:
                progress_callback(i, len(sheet_names), sheet)

            try:
                reader = RawSheetReader(wb_path, sheet)
            except Exception as e:
                log.warning("Error leyendo hoja '%s': %s", sheet, e)
                continue

            structure = detect_structure(reader)
            base_name = _sanitize(f"{wb_path.stem}__{sheet}")[:48]

            if structure == "timeseries":
                try:
                    geo_dfs = parse_sheet_timeseries(wb_path, sheet)
                except Exception as e:
                    log.warning("Error parseando timeseries '%s': %s", sheet, e)
                    geo_dfs = {}

                for geo_level, df in geo_dfs.items():
                    if df is None or df.empty:
                        continue
                    tname = f"{base_name}__{geo_level}"[:64]
                    df.to_sql(tname, conn, if_exists="replace", index=False)
                    manifest.append({
                        "archivo": wb_path.name,
                        "hoja": sheet,
                        "tabla": tname,
                        "estructura": f"timeseries_{geo_level}",
                        "filas": len(df),
                        "columnas": len(df.columns),
                    })
                    log.info("  ✓ %s → %s (%d f)", sheet, tname, len(df))

            else:  # flat
                try:
                    df = parse_sheet_flat(wb_path, sheet)
                except Exception as e:
                    log.warning("Error parseando flat '%s': %s", sheet, e)
                    df = None

                if df is not None and not df.empty:
                    tname = base_name[:64]
                    df.to_sql(tname, conn, if_exists="replace", index=False)
                    manifest.append({
                        "archivo": wb_path.name,
                        "hoja": sheet,
                        "tabla": tname,
                        "estructura": "flat",
                        "filas": len(df),
                        "columnas": len(df.columns),
                    })
                    log.info("  ✓ %s → %s (%d f)", sheet, tname, len(df))

    return manifest
